"""Deterministic stand-in for an external DOCX editor (qualification only).

The foreign lane exists because some edits happen OUTSIDE this engine.  A
qualification case therefore needs a step that edits a candidate the way an
external tool does: either byte surgery on one package part for a text probe,
or a real DOCX package edit for the bounded media-add probe.  Every operation
leaves the engine's typed writer out of the mutation path.

It is deliberately not wired into any product surface: it is the qualification
harness's fake mutator, the same role the in-test helpers play, kept here so
the frozen plan can invoke it through the ordinary CLI adapter.
"""
from __future__ import annotations

import argparse
import base64
import io
import sys
import zipfile
from pathlib import Path

from docx import Document

_MEDIA_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC"
)


def rewrite_document_xml(candidate: Path, old: str, new: str, *, part: str = "word/document.xml") -> int:
    """Replace ``old`` with ``new`` in one part; return the replacement count.

    Every other member is copied byte-for-byte, and the member order is
    preserved, so the only difference from the input package is the declared
    substitution.  An absent ``old`` is a hard error: a case that silently
    edited nothing would pass for the wrong reason."""
    source = zipfile.ZipFile(candidate)
    buffer = io.BytesIO()
    count = 0
    try:
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == part:
                    text = data.decode("utf-8")
                    count = text.count(old)
                    if count == 0:
                        raise SystemExit(f"external-edit: {old!r} is absent from {part}")
                    data = text.replace(old, new).encode("utf-8")
                target.writestr(item, data)
    finally:
        source.close()
    candidate.write_bytes(buffer.getvalue())
    return count


def add_media(candidate: Path) -> None:
    """Add one PNG drawing to the first document paragraph."""
    document = Document(str(candidate))
    if not document.paragraphs:
        raise SystemExit("external-edit: candidate has no document paragraph")
    document.paragraphs[0].add_run().add_picture(io.BytesIO(_MEDIA_PNG))
    document.save(str(candidate))



def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("candidate", help="candidate .docx to edit in place")
    parser.add_argument("old", nargs="?", help="exact text to replace")
    parser.add_argument("new", nargs="?", help="replacement text")
    parser.add_argument("--part", default="word/document.xml")
    parser.add_argument("--media-add", action="store_true")
    args = parser.parse_args(argv)
    if args.media_add:
        if args.old is not None or args.new is not None:
            parser.error("--media-add does not accept old/new text arguments")
        add_media(Path(args.candidate))
        print("external-edit: added one PNG drawing")
        return 0
    if args.old is None or args.new is None:
        parser.error("old and new are required unless --media-add is used")
    count = rewrite_document_xml(Path(args.candidate), args.old, args.new, part=args.part)
    print(f"external-edit: replaced {count} occurrence(s)")
    return 0


def _self_check() -> None:
    import tempfile

    folder = Path(tempfile.mkdtemp(prefix="external-edit-"))
    candidate = folder / "candidate.docx"
    with zipfile.ZipFile(candidate, "w") as archive:
        archive.writestr("word/document.xml", "<w:t>A</w:t><w:t>末</w:t>")
        archive.writestr("word/styles.xml", "<w:styles/>")
    assert rewrite_document_xml(candidate, "A", "B") == 1
    with zipfile.ZipFile(candidate) as archive:
        assert archive.read("word/document.xml").decode("utf-8") == "<w:t>B</w:t><w:t>末</w:t>"
        # untouched members survive byte-for-byte
        assert archive.read("word/styles.xml") == b"<w:styles/>"
    try:
        rewrite_document_xml(candidate, "absent", "x")
    except SystemExit:
        pass
    else:
        raise AssertionError("an absent anchor must refuse rather than no-op")
    print("external-edit self-check: ok")


if __name__ == "__main__":
    if len(sys.argv) == 2 and sys.argv[1] == "--self-check":
        _self_check()
    else:
        raise SystemExit(main())
