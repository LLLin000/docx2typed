"""docx2typed: vertical alignment as text (schema 2).

The contract (docs/prd/vertical-alignment-as-text.md): superscript/subscript is
a textual annotation of editable content **iff** the engine can factor
`w:vertAlign` out of the run formatting reversibly and losslessly. Eligibility
is canonical run-property identity — remove exactly one `w:vertAlign` and the
remainder must canonicalise equal to the paragraph's base style — never a
feature-dict diff, because `rpr_features` enumerates only part of Word's `rPr`
vocabulary.
"""
from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from docx import Document

from scripts import main
from scripts.mcp_server import (
    build_docx,
    batch_edit,
    commit_sync,
    document_patch,
    document_search,
    format_span,
    session,
    workdir_open,
)
from scripts.typed_core import (
    NS_W,
    Paragraph,
    StyleRegistry,
    TextNode,
    TypedError,
    content_signature,
    factor_vertical_style,
    merge_adjacent_text,
    parse_typed,
    promote_vertical_alignment,
    serialize_typed,
    skeleton,
    vertical_style_variant,
)
from scripts.edit import refresh_edit_projection
from scripts.typed_docx import (
    build_workdir,
    extract_workdir,
    json_bytes,
    parse_package_document,
    ValidationError,
    sha256_file,
    validate_workdir,
)

HEAD = (
    '<!--@typed schema="{schema}" format="format.json" styles="styles.json" '
    'template="_template.docx" source="x.docx"-->\n'
)



def _code(result) -> str:
    envelope = _envelope(result)
    return (envelope.get("diagnostics") or [{}])[0].get("code") or "OK"


def _envelope(result):
    if isinstance(result, str):
        return json.loads(result)
    if isinstance(result, dict):
        return result
    return result.structuredContent


def _registry(*rprs: str) -> StyleRegistry:
    registry = StyleRegistry()
    for rpr in rprs:
        registry.ensure(rpr)
    return registry


def _rpr(*children: str) -> str:
    return f'<w:rPr xmlns:w="{NS_W}">' + "".join(children) + "</w:rPr>"


BODY_FONT = '<w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/>'
SUPER = '<w:vertAlign w:val="superscript"/>'
SUB = '<w:vertAlign w:val="subscript"/>'


# ---------------------------------------------------------------------------
# The predicate: canonical identity, one element removed
# ---------------------------------------------------------------------------

def test_factors_a_pure_vertical_variant():
    registry = _registry(_rpr(BODY_FONT), _rpr(BODY_FONT, SUPER))
    base = registry.ensure(_rpr(BODY_FONT))
    variant = registry.ensure(_rpr(BODY_FONT, SUPER))
    assert factor_vertical_style(registry, variant, base) == ("superscript", base)


def test_refuses_when_something_else_differs_too():
    registry = _registry(_rpr(BODY_FONT), _rpr(BODY_FONT, SUPER), _rpr('<w:rFonts w:ascii="宋体"/>', SUPER))
    base = registry.ensure(_rpr(BODY_FONT))
    hinted = registry.ensure(_rpr('<w:rFonts w:ascii="宋体"/>', SUPER))
    # the font differs as well: textifying would silently drop it
    assert factor_vertical_style(registry, hinted, base) is None


def test_refuses_position_and_non_alignment_values():
    registry = _registry(
        _rpr(BODY_FONT),
        _rpr(BODY_FONT, '<w:position w:val="6"/>'),
        _rpr(BODY_FONT + '<w:vertAlign w:val="baseline"/>'),
    )
    base = registry.ensure(_rpr(BODY_FONT))
    positioned = registry.ensure(_rpr(BODY_FONT, '<w:position w:val="6"/>'))
    baseline = registry.ensure(_rpr(BODY_FONT, '<w:vertAlign w:val="baseline"/>'))
    assert factor_vertical_style(registry, positioned, base) is None
    assert factor_vertical_style(registry, baseline, base) is None


def test_refuses_two_vertalign_elements_and_the_base_itself():
    registry = _registry(_rpr(BODY_FONT), _rpr(BODY_FONT, SUPER, SUPER))
    base = registry.ensure(_rpr(BODY_FONT))
    doubled = registry.ensure(_rpr(BODY_FONT, SUPER, SUPER))
    assert factor_vertical_style(registry, doubled, base) is None
    assert factor_vertical_style(registry, base, base) is None


def test_variant_writer_and_predicate_are_inverses():
    """The writer creates the variant the predicate factors back — that is what
    keeps a promoted document byte-identical."""
    registry = _registry(_rpr(BODY_FONT))
    base = registry.ensure(_rpr(BODY_FONT))
    variant = vertical_style_variant(registry, base, "superscript")
    assert (variant, factor_vertical_style(registry, variant, base)) == (
        variant,
        ("superscript", base),
    )


# ---------------------------------------------------------------------------
# Schema 2: the text carries the alignment
# ---------------------------------------------------------------------------

def test_schema2_reads_tags_into_the_vertical_dimension():
    document = parse_typed(
        HEAD.format(schema="2")
        + '\n<!--@p id="P0" base="s_b"-->\n'
        + "Ca^{2+} 与 Ca_{3} 以及字面量 \\^{x}\n"
    )
    nodes = document.paragraphs[0].nodes
    assert [(node.text, node.vertical) for node in nodes] == [
        ("Ca", None),
        ("2+", "superscript"),
        (" 与 Ca", None),
        ("3", "subscript"),
        (" 以及字面量 ^{x}", None),
    ]


def test_schema2_serializes_the_dimension_back_to_tags():
    source = (
        HEAD.format(schema="2")
        + '\n<!--@p id="P0" base="s_b"-->\n'
        + "Ca^{2+} 与 Ca_{3}\n"
    )
    assert serialize_typed(parse_typed(source)) == source


def test_vertical_tag_escapes_braces_and_backslashes_round_trip():
    source = (
        HEAD.format(schema="2")
        + '\n<!--@p id="P0" base="s_b"-->\n'
        + "x^{a\\}b\\\\c}\n"
    )
    document = parse_typed(source)
    assert [(node.text, node.vertical) for node in document.paragraphs[0].nodes] == [
        ("x", None),
        ("a}b\\c", "superscript"),
    ]
    assert serialize_typed(document) == source


def test_schema1_keeps_the_markers_literal():
    source = (
        HEAD.format(schema="1")
        + '\n<!--@p id="P0" base="s_b"-->\n'
        + "Ca^{2+}\n"
    )
    document = parse_typed(source)
    assert [(node.text, node.vertical) for node in document.paragraphs[0].nodes] == [
        ("Ca^{2+}", None)
    ]
    # a schema-1 source keeps its bytes: no escaping, no tags
    assert serialize_typed(document) == source


def test_unknown_schema_is_refused():
    with pytest.raises(TypedError, match="incompatible typed source schema"):
        parse_typed(HEAD.format(schema="3") + '\n<!--@p id="P0" base="s_b"-->\nx\n')


def test_nested_or_unclosed_tags_are_refused():
    for body in ("Ca^{2+", "Ca^{2^{3}}", "Ca^{}"):
        with pytest.raises(TypedError):
            parse_typed(HEAD.format(schema="2") + '\n<!--@p id="P0" base="s_b"-->\n' + body + '\n')


def test_merging_never_crosses_the_vertical_dimension():
    merged = merge_adjacent_text(
        [TextNode("s_b", "Ca"), TextNode("s_b", "2+", "superscript"), TextNode("s_b", " 离子")]
    )
    assert [(node.text, node.vertical) for node in merged] == [
        ("Ca", None),
        ("2+", "superscript"),
        (" 离子", None),
    ]


# ---------------------------------------------------------------------------
# Promotion
# ---------------------------------------------------------------------------

def test_promotion_moves_eligible_regions_only():
    registry = _registry(
        _rpr(BODY_FONT),
        _rpr(BODY_FONT, SUPER),
        _rpr('<w:rFonts w:ascii="宋体"/>', SUPER),
    )
    base = registry.ensure(_rpr(BODY_FONT))
    pure = registry.ensure(_rpr(BODY_FONT, SUPER))
    hinted = registry.ensure(_rpr('<w:rFonts w:ascii="宋体"/>', SUPER))
    document = parse_typed(
        HEAD.format(schema="1")
        + f'\n<!--@p id="P0" base="{base}"-->\n'
        + f'Ca<span data-s="{pure}">2+</span><span data-s="{hinted}">3+</span>\n'
    )
    promoted = promote_vertical_alignment(document, registry)
    assert promoted == 1
    kinds = [
        (node.text, node.style_id, node.vertical)
        for node in document.paragraphs[0].nodes
        if isinstance(node, TextNode)
    ]
    assert kinds == [("Ca", base, None), ("2+", base, "superscript"), ("3+", hinted, None)]


# ---------------------------------------------------------------------------
# End to end: extract promotes, the package keeps the property, byte-identical
# ---------------------------------------------------------------------------

def _superscript_document(path: Path) -> Path:
    document = Document()
    document.add_paragraph("前言。")
    paragraph = document.add_paragraph()
    paragraph.add_run("螯合的Ca")
    superscript = paragraph.add_run("2+")
    superscript.font.superscript = True
    paragraph.add_run("与Cu")
    subscript = paragraph.add_run("2+")
    subscript.font.subscript = True
    document.save(str(path))
    return path


def test_extract_promotes_and_the_built_package_still_carries_the_alignment(tmp_path):
    source = _superscript_document(tmp_path / "s.docx")
    workdir = tmp_path / "wd"
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", "e1"]) == 0

    typed = (workdir / "typed.md").read_text(encoding="utf-8")
    promotion = json.loads((workdir / "format.json").read_text(encoding="utf-8"))["vertical_promotion"]
    assert typed.splitlines()[0].startswith('<!--@typed schema="2"')
    assert "Ca^{2+}" in typed and "Cu_{2+}" in typed
    assert promotion["promoted"] == 2
    validate_workdir(workdir)
    session.workdir = None
    workdir_open(str(workdir), track=False)
    assert commit_sync(label="提升") is not None
    output = tmp_path / "built.docx"
    build_docx(output=str(output))

    import zipfile

    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert 'w:val="superscript"' in xml and 'w:val="subscript"' in xml
    # the markers themselves must never reach the package
    assert "^{" not in xml and "_{" not in xml



def test_refresh_migrates_schema1_without_changing_template(tmp_path):
    source = _superscript_document(tmp_path / "legacy.docx")
    workdir = tmp_path / "legacy-wd"
    extract_workdir(source, workdir)
    template_before = (workdir / "_template.docx").read_bytes()

    typed = parse_typed((workdir / "typed.md").read_text(encoding="utf-8"))
    styles = StyleRegistry.from_json(json.loads((workdir / "styles.json").read_text(encoding="utf-8")))

    def downgrade(nodes, base_style):
        for node in nodes:
            if isinstance(node, TextNode) and node.vertical:
                node.style_id = vertical_style_variant(styles, base_style, node.vertical)
                node.vertical = None
            elif hasattr(node, "children"):
                downgrade(node.children, base_style)

    for paragraph in typed.paragraphs:
        downgrade(paragraph.nodes, paragraph.base_style)
    typed.meta["schema"] = "1"
    (workdir / "typed.md").write_text(serialize_typed(typed), encoding="utf-8", newline="\n")
    styles_text = json_bytes(styles.to_json())
    (workdir / "styles.json").write_bytes(styles_text)

    format_data = json.loads((workdir / "format.json").read_text(encoding="utf-8"))
    format_data.pop("vertical_promotion", None)
    format_data["styles_sha256"] = sha256_file(workdir / "styles.json")
    with zipfile.ZipFile(workdir / "_template.docx") as archive:
        baseline = parse_package_document(archive).document
    for paragraph, record in zip(baseline.paragraphs, format_data["paragraphs"]):
        paragraph.paragraph_id = record["id"]
        record["skeleton"] = skeleton(paragraph.nodes)
    (workdir / "format.json").write_text(
        json.dumps(format_data, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    for name in ("edit.md", "edit.state.json", "edit.state.json.run.json"):
        (workdir / name).unlink(missing_ok=True)

    validate_workdir(workdir)
    refresh_edit_projection(workdir, init=True)

    assert (workdir / "typed.md").read_text(encoding="utf-8").splitlines()[0].startswith(
        '<!--@typed schema="2"'
    )
    assert "Ca^{2+}" in (workdir / "edit.md").read_text(encoding="utf-8")
    assert "Cu_{2+}" in (workdir / "edit.md").read_text(encoding="utf-8")
    evidence = json.loads((workdir / "edit.state.json.run.json").read_text(encoding="utf-8"))
    assert any("vertical-migration schema1->schema2" in item for item in evidence["diagnostics"])
    assert (workdir / "_template.docx").read_bytes() == template_before
    validate_workdir(workdir)

    output = tmp_path / "legacy-built.docx"
    build_workdir(workdir, output)
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output) as built:
        assert built.read("word/document.xml") == original.read("word/document.xml")

def test_schema2_projection_and_search_keep_vertical_annotation(tmp_path):
    source = _superscript_document(tmp_path / "projection.docx")
    workdir = tmp_path / "projection-wd"
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", "projection-extract"]) == 0

    projection = (workdir / "edit.md").read_text(encoding="utf-8")
    assert "Ca^{2+}" in projection and "Cu_{2+}" in projection

    session.workdir = None
    workdir_open(str(workdir), track=False)
    search = json.loads(document_search("Ca2+"))
    assert search["matches"][0]["matched_text"] == "Ca^{2+}"
    assert search["matches"][0]["occurrences"][0]["matched_text"] == "Ca^{2+}"

    (workdir / "edit.md").write_text(
        projection.replace("Ca^{2+}", "Ca^{3+}", 1),
        encoding="utf-8",
    )
    assert _code(commit_sync(operation_id="projection-commit")) == "OK"
    typed = (workdir / "typed.md").read_text(encoding="utf-8")
    assert "螯合的Ca^{3+}与Cu_{2+}" in typed


def test_batch_edit_preserves_vertical_annotation_and_targets_it_explicitly(tmp_path):
    source = _superscript_document(tmp_path / "batch-vertical.docx")
    workdir = tmp_path / "batch-vertical-wd"
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", "batch-extract"]) == 0

    session.workdir = None
    workdir_open(str(workdir), track=False)
    result = batch_edit(
        "P1",
        [{"text": "2+", "vertical": "superscript", "new": "3+"}],
        operation_id="batch-vertical-edit",
    )
    assert _code(result) == "OK"
    typed = (workdir / "typed.md").read_text(encoding="utf-8")
    assert "螯合的Ca^{3+}与Cu_{2+}" in typed


def test_governed_segments_reject_vertical_only_projection_drift(tmp_path):
    source = _superscript_document(tmp_path / "governed-vertical.docx")
    workdir = tmp_path / "governed-vertical-wd"
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", "governed-extract"]) == 0

    session.workdir = None
    workdir_open(str(workdir), track=False)
    projection = (workdir / "edit.md").read_text(encoding="utf-8")
    (workdir / "edit.md").write_text(
        projection.replace("Ca^{2+}", "Ca^{3+}", 1),
        encoding="utf-8",
    )
    assert _code(commit_sync(operation_id="governed-commit")) == "OK"
    format_data = json.loads((workdir / "format.json").read_text(encoding="utf-8"))
    record = next(item for item in format_data["paragraphs"] if item["id"] == "P1")
    assert all(len(segment) == 3 for segment in record["sync_segments"])

    typed = parse_typed((workdir / "typed.md").read_text(encoding="utf-8"))
    target = next(
        node
        for node in typed.paragraphs[1].nodes
        if isinstance(node, TextNode) and node.vertical == "superscript"
    )
    target.vertical = "subscript"
    (workdir / "typed.md").write_text(serialize_typed(typed), encoding="utf-8", newline="\n")

    with pytest.raises(ValidationError, match="governed-formatting-changed"):
        validate_workdir(workdir)



# ---------------------------------------------------------------------------
# The signature must see the dimension (a vertical-only change is a change)
# ---------------------------------------------------------------------------

def test_content_signature_sees_a_vertical_only_change():
    plain = Paragraph("P0", "s_b", [TextNode("s_b", "Ca2+")])
    superscript = Paragraph(
        "P0", "s_b", [TextNode("s_b", "Ca2+", "superscript")]
    )
    assert content_signature(plain) != content_signature(superscript)


def test_vertical_only_format_is_built_not_replayed(tmp_path):
    """A change that only adds the alignment, leaving every character alone, must
    reach the package: if the content signature ignored the dimension the build
    would treat the paragraph as untouched and replay the baseline bytes."""
    import zipfile

    source = tmp_path / "s.docx"
    document = Document()
    document.add_paragraph("第一段正文。")
    variant = document.add_paragraph()
    variant.add_run("2+").font.superscript = True  # gives the doc a superscript style
    document.save(str(source))
    workdir = tmp_path / "wd"
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", "e1"]) == 0

    session.workdir = None
    workdir_open(str(workdir), track=False)
    assert _code(format_span(
        paragraph_id="P0", old="第一段正文。", attributes={"vertAlign": "superscript"}
    )) == "OK"
    assert _code(commit_sync(label="仅上标")) == "OK"
    output = tmp_path / "built.docx"
    assert _code(build_docx(output=str(output))) == "OK"

    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    i = xml.find("第一段正文。")
    start = max(xml.rfind("<w:p>", 0, i), xml.rfind("<w:p ", 0, i))
    paragraph = xml[start : xml.find("</w:p>", i)]
    assert 'w:val="superscript"' in paragraph, paragraph[-260:]


def test_baseline_format_clears_canonical_vertical_dimension(tmp_path):
    source = _superscript_document(tmp_path / "baseline.docx")
    workdir = tmp_path / "baseline-wd"
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", "baseline-extract"]) == 0

    session.workdir = None
    workdir_open(str(workdir), track=False)
    assert _code(format_span(
        paragraph_id="P1",
        old="Cu2+",
        attributes={"vertAlign": "baseline"},
        operation_id="baseline-format",
    )) == "OK"

    typed = (workdir / "typed.md").read_text(encoding="utf-8")
    assert "Ca^{2+}与Cu2+" in typed


def test_document_patch_writes_vertical_tags_canonically(tmp_path):
    source = _superscript_document(tmp_path / "patch.docx")
    workdir = tmp_path / "patch-wd"
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", "patch-extract"]) == 0

    session.workdir = None
    workdir_open(str(workdir), track=False)
    assert not document_patch(
        hunks=[{"paragraph_id": "P1", "old": "与Cu2+", "new": "与Cu^{3+" + "}"}],
        operation_id="patch-vertical",
    ).isError
    assert _code(commit_sync(operation_id="patch-vertical-save")) == "OK"

    typed = (workdir / "typed.md").read_text(encoding="utf-8")
    assert "螯合的Ca^{2+}与Cu^{3+}" in typed
