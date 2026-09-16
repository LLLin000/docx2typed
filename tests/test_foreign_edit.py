"""Receipt-bound foreign candidate admission with a deterministic fake mutator."""
from __future__ import annotations

import base64
import io
import json
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document

from scripts.protocol import semantic_sha256

from scripts import main
from scripts.mcp_server import (
    _commit_decision,
    build_docx,
    commit_sync,
    foreign_edit_adopt,
    foreign_edit_prepare,
    replace_text,
    session,
    workspace_identity_fields,
    workdir_open,
)
from scripts.foreign_edit import (
    ForeignEditError,
    build_provenance,
    declared_tool_provenance,
    make_receipt,
    receipt_bytes,
    store_receipt_path,
    validate_provenance,
)
from scripts.store import head_version, history_list


def _data(result) -> dict:
    payload = result.structuredContent
    return payload.get("data", {}) if isinstance(payload, dict) else {}


def _failure(result) -> str | None:
    if not getattr(result, "isError", False):
        return None
    payload = result.structuredContent or {}
    return (payload.get("diagnostics") or [{}])[0].get("code")


def _make_docx(
    path: Path,
    *,
    include_media: bool = False,
    first_text: str = "Initial paragraph one",
) -> None:
    document = Document()
    document.add_paragraph(first_text)
    document.add_paragraph("Base paragraph two")
    if include_media:
        document.paragraphs[0].add_run().add_picture(io.BytesIO(_MEDIA_PNG))
    document.save(str(path))


def _open(
    tmp_path: Path,
    name: str,
    *,
    include_media: bool = False,
    commit_p1: bool = False,
) -> Path:
    source = tmp_path / f"{name}.docx"
    _make_docx(source, include_media=include_media, first_text="Base paragraph one" if commit_p1 else "Initial paragraph one")
    workdir = tmp_path / name
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", f"{name}-extract"]) == 0
    session.workdir = None
    session.last_build_output = None
    assert _failure(workdir_open(str(workdir), track=False)) is None
    if commit_p1:
        assert _failure(replace_text("P1", "Base paragraph two", "Native paragraph two", operation_id=f"{name}-edit")) is None
    else:
        assert _failure(replace_text("P0", "Initial paragraph one", "Base paragraph one", operation_id=f"{name}-edit")) is None
    saved = commit_sync(operation_id=f"{name}-save")
    assert _failure(saved) is None, saved
    assert _data(saved).get("version", {}).get("created"), json.dumps(_data(saved))
    assert not _commit_decision(workdir)["version_dirty"], json.dumps(_commit_decision(workdir))
    return workdir


def _mutate_text(candidate: Path, old: str, new: str) -> None:
    temporary = candidate.with_suffix(".mutated.docx")
    with zipfile.ZipFile(candidate, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml":
                assert old.encode("utf-8") in payload
                payload = payload.replace(old.encode("utf-8"), new.encode("utf-8"), 1)
            target.writestr(info, payload)
    temporary.replace(candidate)


def _mutate_format(candidate: Path) -> None:
    temporary = candidate.with_suffix(".format.docx")
    with zipfile.ZipFile(candidate, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml":
                marker = b"<w:r>"
                assert marker in payload
                payload = payload.replace(
                    marker,
                    b"<w:r><w:rPr><w:b/></w:rPr>",
                    1,
                )
            target.writestr(info, payload)
    temporary.replace(candidate)

def _mutate_superscript(candidate: Path) -> None:
    """Split a plain ``Ca2+`` run and superscript only ``2+``."""
    temporary = candidate.with_suffix(".vertical.docx")
    old = (
        b'<w:r><w:rPr xmlns:w="http://schemas.openxmlformats.org/'
        b'wordprocessingml/2006/main" /><w:t>Ca2+</w:t></w:r>'
    )
    new = (
        b'<w:r><w:rPr xmlns:w="http://schemas.openxmlformats.org/'
        b'wordprocessingml/2006/main" /><w:t>Ca</w:t></w:r>'
        b'<w:r><w:rPr xmlns:w="http://schemas.openxmlformats.org/'
        b'wordprocessingml/2006/main"><w:vertAlign w:val="superscript"/>'
        b'</w:rPr><w:t>2+</w:t></w:r>'
    )
    with zipfile.ZipFile(candidate, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml":
                assert old in payload
                payload = payload.replace(old, new, 1)
            target.writestr(info, payload)
    temporary.replace(candidate)

_MEDIA_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)
_ADDED_MEDIA_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC"
)


def _add_media(candidate: Path, paragraph_index: int = 0) -> None:
    document = Document(str(candidate))
    document.paragraphs[paragraph_index].add_run().add_picture(io.BytesIO(_ADDED_MEDIA_PNG))
    document.save(str(candidate))


def _mutate_media_package(candidate: Path, mutation: str) -> None:
    assert mutation in {"media", "relationship-target", "rid-reuse", "lost-part"}
    temporary = candidate.with_suffix(f".{mutation}.docx")
    with zipfile.ZipFile(candidate, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        removed = False
        for info in source.infolist():
            if mutation == "lost-part" and info.filename == "word/settings.xml":
                removed = True
                continue
            payload = source.read(info.filename)
            if mutation == "media" and info.filename == "word/media/image1.png":
                assert payload
                payload = payload[:-1] + bytes([payload[-1] ^ 1])
            elif mutation == "relationship-target" and info.filename == "word/_rels/document.xml.rels":
                marker = b'Target="media/image1.png"'
                assert marker in payload
                payload = payload.replace(marker, b'Target="media/image2.png"', 1)
            elif mutation == "rid-reuse" and info.filename == "word/_rels/document.xml.rels":
                match = re.search(
                    rb'<Relationship Id="([^"]+)"[^>]*Target="media/image2\.png"',
                    payload,
                )
                assert match
                payload = payload.replace(
                    b'Id="' + match.group(1) + b'"',
                    b'Id="rId9"',
                    1,
                )
            target.writestr(info, payload)
    assert mutation != "lost-part" or removed
    temporary.replace(candidate)


def test_target_owned_media_addition_is_adopted_and_round_trips(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "media-positive")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="media-prepare")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _add_media(candidate)

    adopted = foreign_edit_adopt(str(candidate), operation_id="media-adopt")
    assert _failure(adopted) is None, adopted
    result = _data(adopted)
    assert any(item["kind"] == "media-add" and item["paragraph_id"] == "P0" for item in result["changes"])
    assert head_version(workdir)["version"] == "V2"

    output = tmp_path / "media-built.docx"
    built = build_docx(output=str(output), operation_id="media-build")
    assert _failure(built) is None, built
    with zipfile.ZipFile(output) as archive:
        assert "word/media/image1.png" in archive.namelist()


@pytest.mark.parametrize("mutation", ["media", "relationship-target", "rid-reuse", "lost-part"])
def test_media_addition_rejects_existing_media_retarget_or_loss(
    tmp_path: Path, mutation: str
) -> None:
    workdir = _open(tmp_path, f"media-{mutation}", include_media=True, commit_p1=True)
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id=f"{mutation}-prepare")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _add_media(candidate)
    _mutate_media_package(candidate, mutation)

    refused = foreign_edit_adopt(str(candidate), operation_id=f"{mutation}-adopt")
    assert _failure(refused) in {"foreign-opaque-or-package-changed", "foreign-package-invalid"}
    assert head_version(workdir)["version"] == "V1"


def test_media_addition_outside_target_is_refused(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "media-scope")
    prepared = foreign_edit_prepare(target=["paragraph:P1"], operation_id="media-scope-prepare")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _add_media(candidate, paragraph_index=0)

    refused = foreign_edit_adopt(str(candidate), operation_id="media-scope-adopt")
    assert _failure(refused) == "foreign-out-of-scope"
    assert head_version(workdir)["version"] == "V1"


def test_media_addition_rejects_unowned_document_change(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "media-unowned")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="media-unowned-prepare")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _add_media(candidate)
    _mutate_unattributed_document(candidate)

    refused = foreign_edit_adopt(str(candidate), operation_id="media-unowned-adopt")
    assert _failure(refused) == "foreign-opaque-or-package-changed"
    assert head_version(workdir)["version"] == "V1"



def test_vertical_only_change_is_not_a_foreign_noop(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "vertical-only")
    assert _failure(
        replace_text("P0", "Base paragraph one", "Ca2+", operation_id="vertical-base-edit")
    ) is None
    assert _failure(commit_sync(operation_id="vertical-base-save")) is None

    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="vertical-prepare")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _mutate_superscript(candidate)

    adopted = foreign_edit_adopt(str(candidate), operation_id="vertical-adopt")
    assert _failure(adopted) is None, adopted
    result = _data(adopted)
    assert any(item["paragraph_id"] == "P0" for item in result["changes"])
    assert result["decision"] == "auto"
    assert head_version(workdir)["version"] == "V3"
    assert "Ca^{2+}" in (workdir / "typed.md").read_text(encoding="utf-8")


def _mutate_style_definition(candidate: Path) -> None:
    temporary = candidate.with_suffix(".style.docx")
    with zipfile.ZipFile(candidate, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/styles.xml":
                marker = b'w:val="Normal"'
                assert marker in payload
                payload = payload.replace(marker, b'w:val="Normal foreign"', 1)
            target.writestr(info, payload)
    temporary.replace(candidate)


def _mutate_normalization(candidate: Path) -> None:
    temporary = candidate.with_suffix(".normalization.docx")
    with zipfile.ZipFile(candidate, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml":
                marker = b"<w:document "
                assert marker in payload
                payload = payload.replace(
                    marker,
                    b'<w:document w:rsidR="00000000" ',
                    1,
                )
            target.writestr(info, payload)
    temporary.replace(candidate)


def _mutate_unattributed_document(candidate: Path) -> None:
    temporary = candidate.with_suffix(".unattributed.docx")
    with zipfile.ZipFile(candidate, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml":
                marker = b"<w:document "
                assert marker in payload
                payload = payload.replace(
                    marker,
                    b'<w:document w:customAttr="foreign" ',
                    1,
                )
            target.writestr(info, payload)
    temporary.replace(candidate)

def _add_opaque_part(candidate: Path) -> None:
    temporary = candidate.with_suffix(".opaque.docx")
    with zipfile.ZipFile(candidate, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        target.writestr("customXml/item2.xml", b"<opaque/>\n")
    temporary.replace(candidate)


def _mutate_section(candidate: Path) -> None:
    temporary = candidate.with_suffix(".section.docx")
    with zipfile.ZipFile(candidate, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml":
                assert b'<w:pgSz w:w="12240" w:h="15840"/>' in payload
                payload = payload.replace(
                    b'<w:pgSz w:w="12240" w:h="15840"/>',
                    b'<w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/>',
                    1,
                )
            target.writestr(info, payload)
    temporary.replace(candidate)


def _append_paragraph(candidate: Path, text: str) -> None:
    """Append one paragraph the way an external editor saves: a new w:p plus
    whatever package churn that tool brings."""
    document = Document(str(candidate))
    document.add_paragraph(text)
    document.save(str(candidate))


def _insert_paragraph_after(candidate: Path, after_index: int, text: str) -> None:
    """Insert a paragraph mid-document, the way an external editor does."""
    document = Document(str(candidate))
    new_paragraph = document.add_paragraph(text)
    document.paragraphs[after_index]._p.addnext(new_paragraph._p)
    document.save(str(candidate))


def test_prepare_rejects_live_workdir_output(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "prepare-isolation")
    result = foreign_edit_prepare(
        output=str(workdir / "candidate.docx"),
        operation_id="prepare-isolation",
    )
    assert _failure(result) == "foreign-candidate-invalid"
    assert not (workdir / "candidate.docx").exists()


def test_prepare_fake_mutator_adopts_next_version_same_workspace(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "prepared")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="prepare-1")
    assert _failure(prepared) is None, prepared
    details = _data(prepared)
    candidate = Path(details["candidate"])
    locator = Path(details["receipt"])
    assert candidate.is_file()
    assert locator.is_file()
    # The sidecar beside the candidate only locates authority...
    locator_data = json.loads(locator.read_text(encoding="utf-8"))
    assert locator_data["schema"] == "docx2typed-foreign-candidate-locator-1"
    candidate_id = str(details["candidate_id"])
    assert locator_data["candidate_id"] == candidate_id
    # ...and the authoritative receipt lives in the engine store.
    stored = json.loads(
        store_receipt_path(workdir, candidate_id).read_text(encoding="utf-8")
    )
    assert stored["schema"] == "docx2typed-foreign-candidate-1"
    assert stored["base_version"] == "V1"

    _mutate_text(candidate, "Base paragraph one", "Foreign paragraph one")
    adopted = foreign_edit_adopt(str(candidate), operation_id="adopt-1")
    assert _failure(adopted) is None, json.dumps(adopted.structuredContent, ensure_ascii=False)
    result = _data(adopted)
    assert result["base_version"] == "V1"
    assert result["decision"] == "auto"
    assert head_version(workdir)["version"] == "V2"
    assert workspace_identity_fields(workdir)["family_id"] == stored["family_id"]
    assert workspace_identity_fields(workdir)["workspace_id"] == stored["workspace_id"]
    version = next(item for item in history_list(workdir)["versions"] if item["version"] == "V2")
    assert version["origin"] == "baseline-transition"
    assert version["metadata"]["foreign_edit"]["candidate_id"] == candidate_id
    assert "Foreign paragraph one" in (workdir / "typed.md").read_text(encoding="utf-8")


def test_out_of_target_and_opaque_changes_refuse_without_moving_head(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "gates")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="prepare-gates")
    candidate = Path(_data(prepared)["candidate"])
    _mutate_text(candidate, "Base paragraph one", "Foreign paragraph one")
    _mutate_text(candidate, "Base paragraph two", "Unexpected paragraph two")
    refused = foreign_edit_adopt(str(candidate), operation_id="adopt-outside")
    assert _failure(refused) == "foreign-out-of-scope"
    assert head_version(workdir)["version"] == "V1"

    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="prepare-opaque")
    opaque_candidate = Path(_data(prepared)["candidate"])
    _add_opaque_part(opaque_candidate)
    refused = foreign_edit_adopt(str(opaque_candidate), operation_id="adopt-opaque")
    assert _failure(refused) in {"foreign-opaque-or-package-changed", "foreign-package-invalid"}
    assert head_version(workdir)["version"] == "V1"


def test_head_drift_and_missing_receipt_require_explicit_lineage(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "lineage")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="prepare-drift")
    assert _failure(prepared) is None
    candidate = Path(_data(prepared)["candidate"])
    _mutate_text(candidate, "Base paragraph one", "Foreign paragraph one")
    assert _failure(replace_text("P1", "Base paragraph two", "Native paragraph two", operation_id="native-1")) is None
    assert _failure(commit_sync(operation_id="native-save")) is None
    assert _failure(foreign_edit_adopt(str(candidate), operation_id="adopt-drift")) == "foreign-edit-conflict"
    assert head_version(workdir)["version"] == "V2"

    assert _failure(workdir_open(str(workdir), track=False)) is None
    # The candidate is deliberately detached from the receipt path.
    manual = tmp_path / "manual-candidate.docx"
    with zipfile.ZipFile(candidate, "r") as source_zip, zipfile.ZipFile(manual, "w") as target_zip:
        for info in source_zip.infolist():
            target_zip.writestr(info, source_zip.read(info.filename))
    manual_result = foreign_edit_adopt(str(manual), operation_id="manual-no-lineage")
    assert _failure(manual_result) == "foreign-lineage-required"


def test_manual_candidate_needs_explicit_family_and_base(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "manual")
    manual = tmp_path / "manual-candidate.docx"
    built = build_docx(output=str(manual), operation_id="manual-build")
    assert _failure(built) is None, built
    _mutate_text(manual, "Base paragraph one", "Manual paragraph one")
    identity = workspace_identity_fields(workdir)
    adopted = foreign_edit_adopt(
        str(manual),
        family_id=str(identity["family_id"]),
        base_version="V1",
        target=["paragraph:P0"],
        operation_id="manual-adopt",
    )
    assert _failure(adopted) is None, adopted
    assert head_version(workdir)["version"] == "V2"
    assert Path(str(manual) + ".foreign-candidate.json").is_file()

    # no tool was claimed and no engine runner exists, so the record says so
    version = next(item for item in history_list(workdir)["versions"] if item["version"] == "V2")
    recorded = version["metadata"]["foreign_edit"]["external_provenance"]
    assert recorded["source"] == "manual"
    assert "tool" not in recorded and "argv_redacted" not in recorded


def test_format_change_is_attributed_inside_target(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "format")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="prepare-format")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _mutate_format(candidate)
    adopted = foreign_edit_adopt(str(candidate), operation_id="adopt-format")
    assert _failure(adopted) is None, adopted
    changes = _data(adopted)["changes"]
    assert any(item["kind"] == "format" and item["paragraph_id"] == "P0" for item in changes)
    assert head_version(workdir)["version"] == "V2"


def test_known_structured_xml_normalization_is_recorded(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "normalization")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="prepare-normalization")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _mutate_normalization(candidate)
    adopted = foreign_edit_adopt(str(candidate), operation_id="adopt-normalization")
    assert _failure(adopted) is None, adopted
    result = _data(adopted)
    assert result["normalization"] == ["word/document.xml"]
    assert result["changes"] == []
    assert head_version(workdir)["version"] == "V2"


def test_style_dependency_outside_paragraph_target_is_refused(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "style-scope")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="prepare-style-scope")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _mutate_style_definition(candidate)
    refused = foreign_edit_adopt(str(candidate), operation_id="adopt-style-scope")
    assert _failure(refused) in {"foreign-out-of-scope", "foreign-opaque-or-package-changed"}
    assert head_version(workdir)["version"] == "V1"


def test_explainable_style_dependency_requires_one_consent(tmp_path: Path, monkeypatch) -> None:
    workdir = _open(tmp_path, "style-consent")
    prepared = foreign_edit_prepare(target=["style:Normal"], operation_id="prepare-style-consent")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _mutate_style_definition(candidate)
    first = foreign_edit_adopt(str(candidate), operation_id="adopt-style-consent")
    assert _failure(first) == "foreign-consent-required"
    details = (first.structuredContent.get("diagnostics") or [{}])[0].get("details") or {}
    token = details.get("consent_token")
    assert isinstance(token, str) and token
    monkeypatch.setattr("scripts.foreign_edit.now_iso", lambda: "2099-01-01T00:00:00+00:00")
    second = foreign_edit_adopt(
        str(candidate),
        consent_token=token,
        operation_id="adopt-style-consent-confirmed",
    )
    assert _failure(second) is None, second
    assert _data(second)["decision"] == "consent"
    assert any(item["kind"] == "dependency" for item in _data(second)["changes"])
    assert head_version(workdir)["version"] == "V2"


def test_manual_consent_token_survives_retry_without_receipt(tmp_path: Path, monkeypatch) -> None:
    workdir = _open(tmp_path, "manual-consent")
    manual = tmp_path / "manual-consent-candidate.docx"
    built = build_docx(output=str(manual), operation_id="manual-consent-build")
    assert _failure(built) is None, built
    _mutate_style_definition(manual)
    identity = workspace_identity_fields(workdir)
    kwargs = {
        "family_id": str(identity["family_id"]),
        "base_version": "V1",
        "target": ["style:Normal"],
    }
    first = foreign_edit_adopt(
        str(manual), operation_id="manual-consent-question", **kwargs
    )
    assert _failure(first) == "foreign-consent-required"
    details = (first.structuredContent.get("diagnostics") or [{}])[0].get("details") or {}
    token = details.get("consent_token")
    assert isinstance(token, str) and token
    monkeypatch.setattr("scripts.foreign_edit.now_iso", lambda: "2099-01-01T00:00:00+00:00")
    second = foreign_edit_adopt(
        str(manual),
        consent_token=token,
        operation_id="manual-consent-confirmed",
        **kwargs,
    )
    assert _failure(second) is None, second
    assert head_version(workdir)["version"] == "V2"


def test_page_setup_change_is_attributed_only_inside_a_section_target(tmp_path: Path) -> None:
    outside = _open(tmp_path, "section-outside")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="prepare-section-outside")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _mutate_section(candidate)
    refused = foreign_edit_adopt(str(candidate), operation_id="adopt-section-outside")
    assert _failure(refused) == "foreign-out-of-scope", refused
    assert head_version(outside)["version"] == "V1"

    workdir = _open(tmp_path, "section-inside")
    prepared = foreign_edit_prepare(target=["section:0"], operation_id="prepare-section")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _mutate_section(candidate)
    adopted = foreign_edit_adopt(str(candidate), operation_id="adopt-section")
    assert _failure(adopted) is None, json.dumps(adopted.structuredContent, ensure_ascii=False)
    result = _data(adopted)
    assert result["decision"] == "auto"
    assert [item["kind"] for item in result["changes"]] == ["section"]
    assert result["changes"][0]["section"] == 0
    assert head_version(workdir)["version"] == "V2"


def test_unattributed_structured_package_change_is_refused(tmp_path: Path) -> None:
    workdir = _open(tmp_path, "unattributed")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="prepare-unattributed")
    assert _failure(prepared) is None, prepared
    candidate = Path(_data(prepared)["candidate"])
    _mutate_unattributed_document(candidate)
    refused = foreign_edit_adopt(str(candidate), operation_id="adopt-unattributed")
    assert _failure(refused) == "foreign-opaque-or-package-changed"
    assert head_version(workdir)["version"] == "V1"


def test_hand_edited_export_returns_as_a_version_in_its_own_timeline(tmp_path: Path) -> None:
    """Issue #95: a hand-edited export re-enters its own family as the next
    Version (a baseline transition), not as a new workspace. An omitted target
    is the whole-document manual re-entry route, and a real hand edit changes
    the paragraph count."""
    workdir = _open(tmp_path, "whole-document")
    manual = tmp_path / "hand-edited.docx"
    built = build_docx(output=str(manual), operation_id="whole-document-build")
    assert _failure(built) is None, built
    _mutate_text(manual, "Base paragraph two", "Hand written paragraph two")
    _append_paragraph(manual, "Hand written addition")

    identity = workspace_identity_fields(workdir)
    adopted = foreign_edit_adopt(
        str(manual),
        family_id=str(identity["family_id"]),
        base_version="V1",
        operation_id="whole-document-adopt",
    )
    assert _failure(adopted) is None, json.dumps(adopted.structuredContent, ensure_ascii=False)
    result = _data(adopted)
    assert result["base_version"] == "V1"
    assert result["decision"] == "auto"
    assert head_version(workdir)["version"] == "V2"

    changes = result["changes"]
    assert any(
        item["kind"] == "semantic" and item["paragraph_id"] == "P1" for item in changes
    ), changes
    assert any(item.get("state") == "added" for item in changes), changes
    typed = (workdir / "typed.md").read_text(encoding="utf-8")
    assert "Hand written paragraph two" in typed
    assert "Hand written addition" in typed

    # same family, same timeline: the transition re-roots the baseline only
    after = workspace_identity_fields(workdir)
    assert after["family_id"] == identity["family_id"]
    assert after["workspace_id"] == identity["workspace_id"]
    versions = history_list(workdir)["versions"]
    version = next(item for item in versions if item["version"] == "V2")
    assert version["origin"] == "baseline-transition"
    assert "V1" in [item["version"] for item in versions]


def test_narrow_target_cannot_rest_on_document_order(tmp_path: Path) -> None:
    """Order pairing aligns leftovers; it does not prove which base paragraph a
    candidate paragraph IS, so it may report a change but never authorize one
    against a narrow target. Whole-document mode is unaffected."""
    workdir = _open(tmp_path, "unproven")
    manual = tmp_path / "unproven-candidate.docx"
    built = build_docx(output=str(manual), operation_id="unproven-build")
    assert _failure(built) is None, built
    # an insertion in the middle renumbers the positional ids that follow it
    _insert_paragraph_after(manual, 0, "hand inserted in the middle")
    _mutate_text(manual, "Base paragraph two", "Hand rewritten paragraph two")

    identity = workspace_identity_fields(workdir)
    narrowed = foreign_edit_adopt(
        str(manual),
        family_id=str(identity["family_id"]),
        base_version="V1",
        target=["paragraph:P1"],
        operation_id="unproven-narrow",
    )
    assert _failure(narrowed) == "foreign-identity-unproven", narrowed
    assert head_version(workdir)["version"] == "V1"

    # the same candidate is admissible when the whole document is the target
    adopted = foreign_edit_adopt(
        str(manual),
        family_id=str(identity["family_id"]),
        base_version="V1",
        operation_id="unproven-whole",
    )
    assert _failure(adopted) is None, json.dumps(adopted.structuredContent, ensure_ascii=False)
    assert head_version(workdir)["version"] == "V2"


def test_forged_sidecar_cannot_widen_scope(tmp_path: Path) -> None:
    """The sidecar beside the candidate is a locator; a whole receipt written
    there — self-consistent or not — is never authority."""
    workdir = _open(tmp_path, "sidecar-forgery")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="forgery-prepare")
    details = prepared.structuredContent["data"]
    candidate = Path(details["candidate"])
    receipt_file = Path(details["receipt"])
    _mutate_text(candidate, "Base paragraph two", "Hand edited outside the target")

    identity = workspace_identity_fields(workdir)
    forged = make_receipt(
        candidate_id=str(details["candidate_id"]),
        family_id=str(identity["family_id"]),
        workspace_id=str(identity["workspace_id"]),
        base_version="V1",
        base_commit="0" * 40,
        base_tree="0" * 40,
        base_export_sha256=str(details["base_export_sha256"]),
        target=[],
        anchors=[],
        candidate_original_sha256=str(details["base_export_sha256"]),
    )
    receipt_file.write_bytes(receipt_bytes(forged))

    # a whole receipt cannot be smuggled in where a locator belongs: adoption
    # refuses outright, the store receipt is never reinterpreted
    refused = foreign_edit_adopt(str(candidate), operation_id="forgery-adopt")
    assert _failure(refused) == "foreign-receipt-invalid", refused
    assert head_version(workdir)["version"] == "V1"

    # and a locator pointing at a candidate id the store never heard of
    # resolves to no receipt at all
    stranger = foreign_edit_adopt(
        str(candidate), candidate_id="FCDEADBEEF", operation_id="forgery-stranger"
    )
    assert _failure(stranger) == "foreign-lineage-required", stranger


def test_store_receipt_tampering_fails_validation(tmp_path: Path) -> None:
    """Even the engine-store copy is self-checked: a widened target that does
    not match its digest is refused as an invalid receipt, not adopted."""
    workdir = _open(tmp_path, "store-tamper")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="tamper-prepare")
    details = prepared.structuredContent["data"]
    candidate = Path(details["candidate"])
    candidate_id = str(details["candidate_id"])
    _mutate_text(candidate, "Base paragraph one", "Hand edited inside the target")
    _mutate_text(candidate, "Base paragraph two", "Hand edited outside the target")

    stored = store_receipt_path(workdir, candidate_id)
    data = json.loads(stored.read_text(encoding="utf-8"))
    data["target"] = []
    stored.write_text(json.dumps(data), encoding="utf-8")

    refused = foreign_edit_adopt(
        str(candidate), candidate_id=candidate_id, operation_id="tamper-adopt"
    )
    assert _failure(refused) == "foreign-receipt-invalid", refused
    assert head_version(workdir)["version"] == "V1"


def test_provenance_redacts_and_binds_engine_digests() -> None:
    """A caller may name a tool and paste an argv; it may not attest a digest
    the engine did not observe, and no path or secret survives redaction."""
    record = declared_tool_provenance(
        tool_name="officecli",
        tool_version="1.2.3",
        argv=[
            "batch",
            r"C:\Users\secret\docs\老师修改.docx",
            "--token=hunter2abc",
            "--out",
            "/srv/tmp/final.docx",
        ],
        exit_code=0,
        receipt_digest="r" * 64,
        input_candidate_sha256="i" * 64,
        output_candidate_sha256="o" * 64,
        analysis_digest="a" * 64,
    )
    joined = " ".join(record["argv_redacted"])
    assert "secret" not in joined and "hunter2abc" not in joined
    assert "老师修改" not in joined and "/srv/tmp" not in joined
    assert record["argv_redacted"][1] == r"<path>"
    assert record["source"] == "declared"
    # the four binding digests are exactly what the engine passed in
    assert record["receipt_digest"] == "r" * 64
    assert record["analysis_digest"] == "a" * 64
    validate_provenance(record)

    # a record that smuggles an un-redacted path back in is refused
    smuggled = dict(record)
    smuggled["argv_redacted"] = [r"C:\Users\secret\docs\a.docx"]
    smuggled["argv_sha256"] = semantic_sha256(smuggled["argv_redacted"])
    with pytest.raises(ForeignEditError) as raised:
        validate_provenance(smuggled)
    assert raised.value.code == "foreign-provenance-invalid"

    # engine-observed is not grantable by a caller: there is no runner
    with pytest.raises(ForeignEditError) as observed:
        build_provenance(
            source="engine-observed",
            receipt_digest="r",
            input_candidate_sha256="i",
            output_candidate_sha256="o",
            analysis_digest="a",
        )
    assert observed.value.code == "foreign-provenance-invalid"


def test_provenance_is_evidence_never_an_admission_input(tmp_path: Path) -> None:
    """The same candidate adopts identically with and without provenance, and
    a malformed provenance record refuses without moving HEAD."""
    workdir = _open(tmp_path, "provenance")
    with_provenance = tmp_path / "prov.docx"
    assert _failure(build_docx(output=str(with_provenance), operation_id="prov-build")) is None
    _mutate_text(with_provenance, "Base paragraph one", "Hand edited paragraph one")
    identity = workspace_identity_fields(workdir)
    accepted = foreign_edit_adopt(
        str(with_provenance),
        family_id=str(identity["family_id"]),
        base_version="V1",
        target=["paragraph:P0"],
        provenance={"tool": "word", "version": "16.0", "argv": ["edit", str(with_provenance)], "exit_code": 0},
        operation_id="prov-adopt",
    )
    assert _failure(accepted) is None, json.dumps(accepted.structuredContent, ensure_ascii=False)
    version = next(item for item in history_list(workdir)["versions"] if item["version"] == "V2")
    recorded = version["metadata"]["foreign_edit"]["external_provenance"]
    assert recorded["source"] == "declared"
    assert recorded["tool"] == {"name": "word", "version": "16.0"}
    assert recorded["argv_redacted"][1] == "<path>"
    assert recorded["analysis_digest"]


def test_provenance_refuses_a_bad_record_without_touching_head(tmp_path: Path) -> None:
    """An otherwise-admissible candidate refuses when its provenance object
    carries a field this engine does not record: the refusal is the
    provenance guard's, HEAD unmoved. (A leaky argv is REDACTED, not refused,
    so only a structurally wrong record reaches this code path.)"""
    workdir = _open(tmp_path, "prov-refuse")
    prepared = foreign_edit_prepare(target=["paragraph:P0"], operation_id="refuse-prep")
    candidate = Path(_data(prepared)["candidate"])
    _mutate_text(candidate, "Base paragraph one", "Hand edited paragraph one")

    rejected = foreign_edit_adopt(
        str(candidate),
        provenance={"tool": "word", "env": {"PATH": "C:\\Users\\nurse"}},
        operation_id="refuse-adopt",
    )
    assert _failure(rejected) == "foreign-provenance-invalid", rejected
    assert head_version(workdir)["version"] == "V1"
    # the refusal names nothing from the rejected record
    assert "nurse" not in json.dumps(rejected.structuredContent, ensure_ascii=False)

