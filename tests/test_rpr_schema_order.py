"""CT_RPr schema-order regression tests for synthesized run properties.

Regression guard for the defect fixed 2026-10-10: ``vertical_style_variant``
appended ``w:vertAlign`` to the end of an rPr. When the base style already
carried an element that must follow ``vertAlign`` in CT_RPr (``w:lang``,
``w:rtl``, ``w:rPrChange``, ...), the synthesized variant was schema-invalid
and Word could reject the document.

The assertions are deliberately three-layered and independent, per review:
  1. element order satisfies the CT_RPr sequence;
  2. attribute semantics survive (size / font / language unchanged);
  3. end-to-end: a real DOCX is still accepted by ``verify``.
"""

import json
import re
from zipfile import ZipFile

import pytest
from docx import Document

from scripts.extract import extract
from scripts.typed_core import (
    _RPR_ORDER,
    _RPR_POSITION,
    Style,
    StyleRegistry,
    insert_rpr_child,
    vertical_base_canonical,
    vertical_style_variant,
)
from scripts.typed_normalize import candidate_report, load_catalog, normalize_workdir
from scripts.verify import verify

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

#: rPr children whose CT_RPr position is AFTER ``w:vertAlign``. A blind append
#: onto a style containing any of these produces invalid XML.
_AFTER_VERTALIGN = tuple(
    name for name in _RPR_ORDER if _RPR_POSITION[name] > _RPR_POSITION["vertAlign"]
)


def _child_order(rpr_xml: str) -> list[str]:
    return [
        name
        for name in re.findall(r"<w:([A-Za-z]+)[ />]", rpr_xml)
        if name != "rPr"
    ]


def _assert_schema_order(rpr_xml: str) -> None:
    """Layer 1: every child appears in CT_RPr sequence."""
    order = _child_order(rpr_xml)
    positions = [_RPR_POSITION[name] for name in order if name in _RPR_POSITION]
    assert positions == sorted(positions), f"rPr children out of schema order: {order}"


def _registry(rpr_cases: dict[str, str]) -> StyleRegistry:
    styles = {}
    for style_id, rpr in rpr_cases.items():
        styles[style_id] = Style(
            style_id=style_id,
            rpr=rpr,
            canonical=style_id,
            label=style_id,
            features={},
            synthesized="",
        )
    return StyleRegistry(styles)


def _rpr(body: str) -> str:
    return f'<w:rPr xmlns:w="{NS_W}">{body}</w:rPr>'


@pytest.mark.parametrize(
    "style_id,body,description",
    [
        (
            "s_szcs",
            '<w:rFonts w:ascii="Arial"/><w:szCs w:val="24"/>',
            "szCs-ending (append happened to be legal here)",
        ),
        (
            "s_lang",
            '<w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/>'
            '<w:sz w:val="21"/><w:szCs w:val="21"/><w:lang w:val="en-US" w:eastAsia="zh-CN"/>',
            "lang-ending (the originally broken case)",
        ),
        (
            "s_rtl",
            '<w:rFonts w:ascii="Arial"/><w:rtl w:val="1"/>',
            "rtl-ending",
        ),
        (
            "s_rprchange",
            '<w:rFonts w:ascii="Arial"/><w:lang w:val="en-US"/><w:rPrChange w:id="1"/>',
            "rPrChange present",
        ),
        (
            "s_plain",
            '<w:rFonts w:hint="eastAsia" w:ascii="Times New Roman"/>',
            "rFonts only",
        ),
    ],
)
def test_synthesized_vertalign_keeps_schema_order(style_id, body, description):
    """Layer 1 + 2: order is legal and the non-vertAlign properties are kept.

    Note: ``rPrChange`` is format history, not a format property — the registry
    canonicalizes it away by design (``strip_rpr_change``), so it is excluded
    from the "properties survive" expectation here. Its *ordering* behaviour in
    the raw insertion primitive is covered separately below.
    """
    registry = _registry({style_id: _rpr(body)})
    variant_id = vertical_style_variant(registry, style_id, "superscript")
    variant_xml = registry.require(variant_id).rpr

    _assert_schema_order(variant_xml)
    assert variant_xml.count("vertAlign") == 1
    # vertAlign must precede every element that CT_RPr places after it
    for name in _AFTER_VERTALIGN:
        if f"<w:{name}" in variant_xml:
            assert variant_xml.index("vertAlign") < variant_xml.index(f"<w:{name}"), (
                f"{description}: w:{name} must not precede w:vertAlign"
            )

    # Layer 2: every original *format property* survives in its original order
    base_order = [n for n in _child_order(_rpr(body)) if n != "rPrChange"]
    variant_order = [n for n in _child_order(variant_xml) if n != "vertAlign"]
    assert variant_order == base_order, "non-vertAlign properties were reordered"


def test_synthesized_vertalign_is_idempotent_and_recognized():
    """Applying the variant twice yields no duplicate; the base is recoverable."""
    style_id = "s_lang"
    body = '<w:rFonts w:ascii="Arial"/><w:sz w:val="21"/><w:lang w:val="en-US"/>'
    base = _rpr(body)
    registry = _registry({style_id: base})

    once = vertical_style_variant(registry, style_id, "superscript")
    # re-synthesizing from an already-variant style returns the same style
    again = vertical_style_variant(registry, once, "superscript")
    assert again == once

    variant_xml = registry.require(once).rpr
    assert variant_xml.count("vertAlign") == 1
    assert vertical_base_canonical(variant_xml) is not None


def test_insert_rpr_child_replaces_in_place_without_duplicates():
    """A same-named child is replaced, never duplicated, and stays ordered."""
    import xml.etree.ElementTree as ET

    ET.register_namespace("w", NS_W)
    root = ET.fromstring(
        _rpr('<w:rFonts w:ascii="Arial"/><w:lang w:val="en-US"/>').encode()
    )
    element = ET.Element(f"{{{NS_W}}}vertAlign", {f"{{{NS_W}}}val": "superscript"})
    insert_rpr_child(root, element)
    insert_rpr_child(root, element)
    out = ET.tostring(root, encoding="unicode")

    assert out.count("vertAlign") == 1
    _assert_schema_order(out)
    assert out.index("vertAlign") < out.index("<w:lang")


def test_insert_rpr_child_places_vertalign_before_rprchange():
    """The primitive keeps schema order when format history is present."""
    import xml.etree.ElementTree as ET

    ET.register_namespace("w", NS_W)
    root = ET.fromstring(
        _rpr(
            '<w:rFonts w:ascii="Arial"/><w:lang w:val="en-US"/>'
            '<w:rPrChange w:id="1"/>'
        ).encode()
    )
    insert_rpr_child(
        root, ET.Element(f"{{{NS_W}}}vertAlign", {f"{{{NS_W}}}val": "superscript"})
    )
    out = ET.tostring(root, encoding="unicode")

    _assert_schema_order(out)
    assert out.index("vertAlign") < out.index("<w:lang")
    assert out.index("vertAlign") < out.index("rPrChange")


def test_vertalign_variant_reaches_a_real_docx(tmp_path):
    """Layer 3: end-to-end, a document using the variant still verifies.

    Exercises the whole lane (extract -> normalize with a superscript
    conversion -> build -> verify) on a style whose rPr ends with ``w:lang``,
    the shape that previously produced invalid XML.
    """
    source = tmp_path / "source.docx"
    workdir = tmp_path / "workdir"
    normalized_docx = tmp_path / "normalized.docx"
    normalized_workdir = tmp_path / "normalized-workdir"

    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Cu²⁺")
    run.font.name = "Times New Roman"
    document.save(source)

    assert extract([str(source), "-o", str(workdir)]) == 0

    candidates = candidate_report(workdir)
    assert candidates, "expected at least one vertical-candidate occurrence"
    format_data = json.loads((workdir / "format.json").read_text(encoding="utf-8"))
    policy = {
        "schema": "vertical-normalization-policy-1",
        "profile": "selective",
        "catalog_hash": load_catalog()["catalog_hash"],
        "template_sha256": format_data["template_sha256"],
        "decisions": {c["occurrence_id"]: "convert" for c in candidates},
    }
    policy_path = tmp_path / "policy.json"
    policy_path.write_text(json.dumps(policy), encoding="utf-8")
    normalize_workdir(workdir, policy_path, normalized_docx, normalized_workdir)

    with ZipFile(normalized_docx) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "w:vertAlign" in xml

    # every rPr that carries vertAlign must be schema-ordered
    for match in re.finditer(r"<w:rPr>.*?</w:rPr>", xml, re.S):
        if "vertAlign" in match.group(0):
            _assert_schema_order(match.group(0))

    assert verify([str(normalized_workdir), str(normalized_docx)]) == 0


def test_rpr_order_constant_places_vertalign_before_lang():
    """Guard the constant itself, so a future edit cannot silently break it."""
    assert _RPR_POSITION["sz"] < _RPR_POSITION["vertAlign"]
    assert _RPR_POSITION["szCs"] < _RPR_POSITION["vertAlign"]
    assert _RPR_POSITION["vertAlign"] < _RPR_POSITION["lang"]
    assert _RPR_POSITION["vertAlign"] < _RPR_POSITION["rtl"]
    assert _RPR_POSITION["vertAlign"] < _RPR_POSITION["rPrChange"]
