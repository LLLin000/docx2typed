"""docx2typed: the insertion chain — an inserted paragraph is ordinary content.

A paragraph created by an insert copies the formatting of the paragraph it was
inserted after, *from the baseline*, which is why the typed state carries
``inherit="Pxx"`` instead of duplicating a style. The editing tools used to see
an empty style there (``unknown style ID: ``, ``protected-context-ambiguous``)
and, worse, a tracked edit inside such a paragraph produced a workdir the
builder refused — but only after the mutation had been written.

The contract under test: reading materialises the inherited style (inserted
content is not special), an insert's content can be formatted and built, and an
edit that would nest tracked changes inside a pending insertion is refused
*before* anything is written.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from scripts import main
from scripts.mcp_server import (
    build_docx,
    commit_sync,
    document_patch,
    format_span,
    get_paragraph,
    insert_paragraph,
    session,
    verify_output,
    workdir_open,
)
from scripts.typed_core import parse_typed, serialize_typed
from scripts.typed_docx import validate_workdir

TYPED_HEAD = (
    '<!--@typed schema="1" format="format.json" styles="styles.json" '
    'template="_template.docx" source="x.docx"-->\n'
)

INSERTED_TEXT = "插入的新段落。"


def _envelope(result):
    if isinstance(result, str):
        return json.loads(result)
    if isinstance(result, dict):
        return result
    return result.structuredContent


def _code(result) -> str:
    envelope = _envelope(result)
    return (envelope.get("diagnostics") or [{}])[0].get("code") or "OK"


def _document(path: Path) -> Path:
    """A document that already carries a superscript style variant — the format
    lane reuses existing variants, so the fixture has to provide one."""
    document = Document()
    document.add_paragraph("第一段正文。")
    paragraph = document.add_paragraph()
    run = paragraph.add_run("上标变体")
    run.font.superscript = True
    document.add_paragraph("第三段正文。")
    document.save(str(path))
    return path


def _workdir(tmp_path: Path) -> Path:
    source = _document(tmp_path / "source.docx")
    workdir = tmp_path / "wd"
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", "e1"]) == 0
    return workdir


def _open(workdir: Path, *, track: bool | None = None) -> None:
    session.workdir = None
    workdir_open(str(workdir), track=track)


def _insert(workdir: Path, after_id: str = "P0") -> str:
    """Insert a paragraph and commit it; returns the paragraph's formal id."""
    assert _code(insert_paragraph(after_id=after_id, text=INSERTED_TEXT)) == "OK"
    assert _code(commit_sync(label="插入一段")) == "OK"
    typed = (workdir / "typed.md").read_text(encoding="utf-8")
    for block in typed.split("<!--@p ")[1:]:
        if INSERTED_TEXT in block:
            return block.split('id="', 1)[1].split('"', 1)[0]
    raise AssertionError("inserted paragraph not found in typed.md")


# ---------------------------------------------------------------------------
# Reading: inserted content is not special
# ---------------------------------------------------------------------------

def test_inherited_style_is_materialized_on_parse():
    document = parse_typed(
        TYPED_HEAD
        + '\n<!--@p id="P0" base="s_abc"-->\n正文\n'
        + '\n<!--@p id="P1" inherit="P0"-->\n插入的一段\n'
    )
    inserted = next(p for p in document.paragraphs if p.paragraph_id == "P1")
    assert inserted.inherit == "P0"
    assert inserted.base_style == "s_abc"
    assert [node.style_id for node in inserted.nodes] == ["s_abc"]
    # the marker still round-trips as an inherit reference
    assert '<!--@p id="P1" inherit="P0"-->' in serialize_typed(document)


def test_inherit_chain_materializes_through_another_insert():
    document = parse_typed(
        TYPED_HEAD
        + '\n<!--@p id="P0" base="s_abc"-->\n正文\n'
        + '\n<!--@p id="P1" inherit="P0"-->\n第一段插入\n'
        + '\n<!--@p id="P2" inherit="P1"-->\n接在插入后面的一段\n'
    )
    chained = next(p for p in document.paragraphs if p.paragraph_id == "P2")
    assert chained.base_style == "s_abc"


def test_unresolvable_inherit_stays_empty_instead_of_guessing():
    document = parse_typed(
        TYPED_HEAD + '\n<!--@p id="P1" inherit="P404"-->\n找不到来源的一段\n'
    )
    assert document.paragraphs[0].inherit == "P404"
    assert document.paragraphs[0].base_style == ""


def test_inserted_paragraph_reports_its_style_to_the_reader(tmp_path):
    workdir = _workdir(tmp_path)
    _open(workdir, track=False)
    paragraph_id = _insert(workdir)
    payload = _envelope(get_paragraph(paragraph_id))
    styles = [entry.get("style_id") for entry in payload.get("styles", [])]
    assert styles and all(styles), payload


# ---------------------------------------------------------------------------
# Writing: the chain works end to end
# ---------------------------------------------------------------------------

def test_inserted_paragraph_can_be_formatted_and_built(tmp_path):
    workdir = _workdir(tmp_path)
    _open(workdir, track=False)
    paragraph_id = _insert(workdir)

    assert _code(format_span(
        paragraph_id=paragraph_id, old=INSERTED_TEXT, attributes={"vertAlign": "superscript"}
    )) == "OK"
    assert _code(commit_sync(label="插入段落上标")) == "OK"
    output = tmp_path / "built.docx"
    assert _code(build_docx(output=str(output))) == "OK"
    assert _code(verify_output(output=str(output))) == "OK"
    validate_workdir(workdir)  # the workdir the edit produced is still valid


def test_nested_revision_is_refused_before_anything_is_written(tmp_path):
    """The old failure mode: the mutation succeeded, then the workdir became
    unbuildable. It must be a refusal instead, with the workdir untouched."""
    workdir = _workdir(tmp_path)
    _open(workdir, track=True)
    paragraph_id = _insert(workdir)
    before = (workdir / "typed.md").read_text(encoding="utf-8")

    result = format_span(
        paragraph_id=paragraph_id, old=INSERTED_TEXT, attributes={"vertAlign": "superscript"}
    )
    assert _code(result) == "edit-inside-pending-insertion"
    assert (workdir / "typed.md").read_text(encoding="utf-8") == before
    validate_workdir(workdir)


def test_patch_into_a_pending_insertion_is_refused_with_the_same_code(tmp_path):
    workdir = _workdir(tmp_path)
    _open(workdir, track=True)
    paragraph_id = _insert(workdir)
    before = (workdir / "typed.md").read_text(encoding="utf-8")

    result = document_patch(
        hunks=[{"paragraph_id": paragraph_id, "old": INSERTED_TEXT, "new": "改过的插入段落。"}]
    )
    assert _code(result) == "edit-inside-pending-insertion"
    assert (workdir / "typed.md").read_text(encoding="utf-8") == before
    validate_workdir(workdir)


def test_a_body_paragraph_is_not_touched_by_the_guard(tmp_path):
    """The control: the guard is about pending insertions, not about track mode."""
    workdir = _workdir(tmp_path)
    _open(workdir, track=True)
    result = format_span(paragraph_id="P0", old="第一段正文。", attributes={"vertAlign": "superscript"})
    assert _code(result) != "edit-inside-pending-insertion"
    validate_workdir(workdir)
